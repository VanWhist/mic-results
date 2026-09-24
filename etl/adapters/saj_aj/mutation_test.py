# -*- coding: utf-8 -*-
"""verify_nc.py の変異テスト（指示書 v2 第4節）。
わざと壊した xlsx / PDF のコピーを作業用フォルダに作り、検証が止めるかを見る。原本は触らない。
使い方: python mutation_test.py <pdf> <xlsx> <rules_json> <golden_json> <workdir> [既知の不具合xlsx:ラベル ...]
"""
import sys, os, json, io, contextlib
import openpyxl
import pypdfium2 as pdfium
from verify_nc import verify, read_pdf, group_rounds, NON_ROUND_SHEETS

PDF, XLSX, RULES, GOLDEN, WORK = sys.argv[1:6]
KNOWN_BAD = [a.rsplit(':', 1) for a in sys.argv[6:]]   # パスにドライブ文字の ':' があるので右から切る
rules = json.load(open(RULES, encoding='utf-8'))
golden = json.load(open(GOLDEN, encoding='utf-8'))
golden_keys = {(g['sheet'], g['bib']) for g in golden['rows']}
os.makedirs(WORK, exist_ok=True)

wb0 = openpyxl.load_workbook(XLSX)
round_sheets = [ws.title for ws in wb0.worksheets if ws.title not in NON_ROUND_SHEETS]
largest = max(round_sheets, key=lambda t: wb0[t].max_row)          # 予選（ページをまたぐことが多い）
final = min(round_sheets, key=lambda t: wb0[t].max_row)            # 最終ラウンド
middle = sorted(round_sheets, key=lambda t: wb0[t].max_row)[len(round_sheets) // 2]


def hdr(ws):
    return [c.value for c in ws[1]]


def col(ws, key):
    return hdr(ws).index(key) + 1


def plain_ok_row(ws, start=2):
    """ゴールデン対象外の完走者の行番号（変異がゴールデン頼みで止まらないようにする）。"""
    for r in range(start, ws.max_row + 1):
        if ws.cell(r, col(ws, 'Status')).value == 'OK' and (ws.title, ws.cell(r, col(ws, 'Bib')).value) not in golden_keys:
            return r
    raise SystemExit('ゴールデン対象外の完走者が見つからない: %s' % ws.title)


def mutate(name, fn):
    wb = openpyxl.load_workbook(XLSX)
    fn(wb)
    path = os.path.join(WORK, name + '.xlsx')
    wb.save(path)
    return path


def m_delete_row(wb):
    ws = wb[largest]; ws.delete_rows(plain_ok_row(ws, 5))


def m_drop_last_page(wb):
    # パーサが最終ページを読み飛ばした状態: その表のラウンドのうち最終ページに載っていた人数ぶん末尾を消す
    year, pages, tables = read_pdf(PDF)
    rounds, _ = group_rounds(tables)
    last = tables[-1]
    rnd = next(r for r in rounds if last in r['tables'])
    ws = wb['%s_%s-%s' % (year, rnd['code'], {'男子': 'm', '女子': 'w'}[rnd['gender']])]
    ws.delete_rows(ws.max_row - len(last['blocks']) + 1, len(last['blocks']))


def m_shift_line2(wb):
    ws = wb[middle]
    keys = ['FIS No', 'Club', 'Jump2', 'DD2', 'Air2 Ja', 'Air2 Jb'] + [h for h in hdr(ws) if str(h).endswith(' Ded')]
    for r in range(ws.max_row, 3, -1):
        for k in keys:
            ws.cell(r, col(ws, k)).value = ws.cell(r - 1, col(ws, k)).value


def base_cols(ws):
    return [h for h in hdr(ws) if str(h).endswith(' Base')]


def m_change_counted(wb):
    """合計に入るジャッジ点を 0.1 動かす。5人制では最高・最低を除くので、中央3人のどれかに当てないと
    合計が変わらず再計算の層が反応しない（2025年のパイロットで判明）。3人制は全員が合計に入る。"""
    ws = wb[middle]
    r = plain_ok_row(ws)
    while True:
        cells = [ws.cell(r, col(ws, h)) for h in base_cols(ws)]
        vals = [c.value for c in cells]
        for c in cells:
            if len(cells) <= 3 or (min(vals) < c.value and round(c.value + 0.1, 1) < max(vals)):
                c.value = round(c.value + 0.1, 1)
                return
        r = plain_ok_row(ws, r + 1)


def m_change_excluded(wb):
    """除外される最高点をさらに 0.1 上げる（合計は変わらない）。止めるのは列位置の照合だけの想定。"""
    ws = wb[middle]
    r = plain_ok_row(ws)
    cells = [ws.cell(r, col(ws, h)) for h in base_cols(ws)]
    top = max(cells, key=lambda c: c.value)
    top.value = round(top.value + 0.1, 1)


def m_swap_turn_judges(wb):
    ws = wb[middle]; r = plain_ok_row(ws)
    for k in range(2, 12):          # 値が違う2列を探して入れ替える（同じ値だと入れ替えにならない）
        a, b = ws.cell(r, col(ws, 'J1 Base')), ws.cell(r, col(ws, 'J2 Base'))
        if a.value != b.value:
            a.value, b.value = b.value, a.value
            return
        r = plain_ok_row(ws, r + 1)


def m_swap_air_judges(wb):
    ws = wb[middle]; r = plain_ok_row(ws)
    while ws.cell(r, col(ws, 'Air1 Ja')).value == ws.cell(r, col(ws, 'Air1 Jb')).value:
        r = plain_ok_row(ws, r + 1)
    a, b = ws.cell(r, col(ws, 'Air1 Ja')), ws.cell(r, col(ws, 'Air1 Jb'))
    a.value, b.value = b.value, a.value


def m_swap_rows(wb):
    ws = wb[largest]; r = plain_ok_row(ws, 10)
    for c in range(1, ws.max_column + 1):
        a, b = ws.cell(r, c), ws.cell(r + 1, c)
        a.value, b.value = b.value, a.value


def m_text_cell(wb):
    ws = wb[final]; r = plain_ok_row(ws)
    ws.cell(r, col(ws, 'J1 Base')).value = '%s %s' % (ws.cell(r, col(ws, 'J1 Base')).value, ws.cell(r, col(ws, 'J2 Base')).value)


def m_wrong_round_code(wb):
    ws = wb[middle]; ws.cell(2, col(ws, 'ラウンド記号')).value = 'SF'


def m_meta_count(wb):
    ws = wb['_meta']; c = ws.cell(2, hdr(ws).index('人数') + 1); c.value = c.value + 1


cases = [
    ('選手の行を1つ消す', mutate('M01_delete_row', m_delete_row), PDF),
    ('最終ページの選手を消す（ページ読み飛ばし）', mutate('M02_drop_last_page', m_drop_last_page), PDF),
    ('2行目を1つ下の選手にずらす', mutate('M03_shift_line2', m_shift_line2), PDF),
    ('合計に入るジャッジ点を1つ 0.1 変える', mutate('M04_change_counted', m_change_counted), PDF),
    ('同じ行のターンジャッジ J1/J2 を入れ替える', mutate('M05_swap_turn_judges', m_swap_turn_judges), PDF),
    ('同じ行のエアジャッジ Ja/Jb を入れ替える', mutate('M06_swap_air_judges', m_swap_air_judges), PDF),
    ('隣り合う2選手の行を入れ替える', mutate('M07_swap_rows', m_swap_rows), PDF),
    ('ジャッジ点2つを1セルの文字列にする', mutate('M08_text_cell', m_text_cell), PDF),
    ('ラウンド記号列を書き換える', mutate('M09_wrong_round_code', m_wrong_round_code), PDF),
    ('_meta の人数を1つ増やす', mutate('M10_meta_count', m_meta_count), PDF),
]
if len(base_cols(wb0[middle])) >= 5:
    cases.append(('除外される最高点を 0.1 上げる（合計は不変、列位置の照合だけが止める想定）',
                  mutate('M11_change_excluded', m_change_excluded), PDF))
for i, exc in enumerate(rules.get('recompute_exceptions', [])):
    # 例外に登録した行でも、別の値が壊れたら止まること（例外が他の壊れ方を覆い隠さない）
    def m_break_exception_row(wb, exc=exc):
        ws = wb[exc['sheet']]
        r = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, col(ws, 'Bib')).value == int(exc['bib']))
        c = ws.cell(r, col(ws, 'Air1 Ja')); c.value = round(c.value + 0.1, 1)
    cases.append(('例外登録した行 BIB %s のエア点を 0.1 変える' % exc['bib'],
                  mutate('M12_exception_row_%d' % i, m_break_exception_row), PDF))
pdf = pdfium.PdfDocument(PDF)
pdf.del_page(len(pdf) - 1)
cut = os.path.join(WORK, 'P01_last_page_removed.pdf')
pdf.save(cut)
cases.append(('PDF のページを1枚消す（xlsx は正しいまま）', XLSX, cut))
for path, label in KNOWN_BAD:
    cases.append(('実在した不具合: ' + label, path, PDF))

stopped, missed = [], []
for label, x, p in cases:
    with contextlib.redirect_stdout(io.StringIO()):
        res = verify(p, x, rules, golden)
    layers = {k: len(v) for k, v in res['errors'].items()}
    (stopped if layers else missed).append((label, layers))
print('対象シート: 最大=%s 中間=%s 最終=%s' % (largest, middle, final))
print('止めた変異:')
for label, layers in stopped:
    print('  OK  %s -> %s' % (label, layers))
print('止められなかった変異:')
for label, layers in missed:
    print('  !!  %s' % label)
if not missed:
    print('  なし')
json.dump(dict(stopped=stopped, missed=missed), open(os.path.join(WORK, 'mutation_result.json'), 'w', encoding='utf-8'),
          ensure_ascii=False, indent=1)
sys.exit(1 if missed else 0)
