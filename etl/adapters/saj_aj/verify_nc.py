# -*- coding: utf-8 -*-
"""全日本MO 国内様式 xlsx の検証（指示書 v2 第4節）。

verify_sajmo.py は parse_sajmo で PDF を読み直して xlsx と比べるので、パーサの読み違いは
読み直しでも同じように起きて一致してしまう。ここでは parse_sajmo も build_sajmo_xlsx も使わず、
PDF の文字の座標から表を組み直し、xlsx（成果物そのもの）と比べる。ラウンド記号とシート名も
ここで独立に決め直して照合する。

層: 完全性 / 印字照合 / 再計算 / 順位 / セル / ゴールデン / 件数ゲート
使い方: python verify_nc.py <pdf> <xlsx> <rules_json> [golden_json]
"""
import sys, re, json, collections
from decimal import Decimal as D, ROUND_DOWN
import pdfplumber
import openpyxl

NUM = re.compile(r'^-?\d+(?:\.\d+)?$')
ROUNDCODE = re.compile(r'^(SF|F|Q)-[a-z/]+$')
STATUS = {'DNF', 'DNS', 'DSQ', 'DQ'}
GENDER_CODE = {'男子': 'm', '女子': 'w'}
CODES_BY_ROUNDS = {1: ['F'], 2: ['Q', 'F'], 3: ['Q', 'F', 'SF']}
ROUND_ORDER = {'Q': 0, 'F': 1, 'SF': 2}
LAYERS = ['完全性', '印字照合', '再計算', '順位', 'セル', 'ゴールデン', '件数ゲート']
NON_ROUND_SHEETS = {'_meta', 'Event Info'}


class NotANumber(ValueError):
    pass


def dec(v):
    if v is None or v == '':
        return None
    if isinstance(v, str) and not NUM.match(v.strip()):
        raise NotANumber(repr(v))
    return D(str(v))


def trunc2(x):
    q = abs(x).quantize(D('0.01'), rounding=ROUND_DOWN)
    return q if x >= 0 else -q


# ---------------------------------------------------------------- PDF 側（パーサ非依存）

def _lines(words, gap=2.5):
    """top の近い語を1行にまとめる。1行目は語ごとに top が 1pt ほどずれて印字される。"""
    out = []
    for w in sorted(words, key=lambda w: w['top']):
        if out and w['top'] - out[-1]['top'] <= gap:
            out[-1]['words'].append(w)
        else:
            out.append({'top': w['top'], 'words': [w]})
    for l in out:
        l['words'].sort(key=lambda w: w['x0'])
    return out


def merge_fragments(row, gap=2.5):
    """同じ行で x の間があいていない語をつなぐ（x0 は先頭の断片のもの）。"""
    out = []
    for w in row:
        if out and w['x0'] - out[-1]['x1'] < gap:
            out[-1] = dict(out[-1], text=out[-1]['text'] + w['text'], x1=w['x1'])
        else:
            out.append(dict(w))
    return out


def column_map(hrow, above):
    """ヘッダの語の x0 から (x0, 1行目の列名, 2行目の列名) を作る。人数は見出しの J の数で決まる。
    数値の多重集合だけでは同じ行の中のジャッジ列の入れ替えが止まらないので、列の位置でも照合する。"""
    cols, n_total, seen_dd, n_turn, n_air = [], 0, False, 0, 0
    for w in hrow:
        t = w['text']
        if re.fullmatch(r'J\d', t):
            if n_total == 0:
                n_turn += 1
                cols.append((w['x0'], 'J%d Base' % n_turn, 'J%d Ded' % n_turn))
            elif seen_dd:
                n_air += 1
                ab = 'Ja' if n_air == 1 else 'Jb'
                cols.append((w['x0'], 'Air1 ' + ab, 'Air2 ' + ab))
        elif t == 'Total':
            n_total += 1
            cols.append((w['x0'], 'Turns Total' if n_total == 1 else 'Air Total', None))
        elif t == 'Jump':
            cols.append((w['x0'], 'Jump1', 'Jump2'))
        elif t == 'DD':
            seen_dd = True
            cols.append((w['x0'], 'DD1', 'DD2'))
        elif t == 'Time':
            cols.append((w['x0'], 'Time', None))
        elif t == 'Point':
            cols.append((w['x0'], 'Time Points', None))
    for w in above:
        if w['text'] == 'スコア':
            cols.append((w['x0'], 'Score', None))
        elif w['text'] == '同点':
            cols.append((w['x0'], 'Tie', None))
    return n_turn, sorted(cols)


def mixed_order(code):
    m = re.match(r'^(?:SF|F|Q)-([a-z/]+)$', code or '')
    if m and '/' in m.group(1):
        return [{'w': '女子', 'm': '男子'}.get(x, '?') for x in m.group(1).split('/')]
    return ['女子', '男子']


def page_pace(words):
    """{性別 or None: ペースタイム}。'23.78 秒' の行か、'ペースタイム' ラベルの直下の数値。
    男女が同じページに並ぶ年は '男子ペースタイム' のように性別つきで印字される。"""
    out = {}
    sec = next((w for w in words if w['text'] == '秒'), None)
    if sec:
        row = [w for w in words if abs(w['top'] - sec['top']) < 2 and NUM.match(w['text'])]
        if row:
            out[None] = D(row[0]['text'])
    for lab in (w for w in words if w['text'].endswith('ペースタイム')):
        # 値の位置は年で違う（2024・2025 はラベルの直下 x+23、2022 は右下 x+115）
        below = sorted([w for w in words if NUM.match(w['text']) and 0 < w['top'] - lab['top'] < 30
                        and -10 < w['x0'] - lab['x0'] < 130], key=lambda w: (w['top'], w['x0']))
        if below:
            g = next((k for k in GENDER_CODE if lab['text'].startswith(k)), None)
            out.setdefault(g, D(below[0]['text']))
    return out


def read_pdf(path):
    """表ごとに {page, code, heading, gender, pace, nturn, columns, blocks} を返す。
    1ページに【女子】【男子】の2表が並ぶ年がある（2024 SF-w/m）。
    選手ブロックは「BIB 列に整数がある行」から次のその行までで、1選手=2行を期待する。"""
    tables, pages = [], []
    with pdfplumber.open(path) as pdf:
        year = None
        for pno, page in enumerate(pdf.pages, 1):
            words = page.extract_words()
            if year is None:
                y = next((re.match(r'(\d{4})年', w['text']) for w in words if re.match(r'\d{4}年', w['text'])), None)
                year = y.group(1) if y else None
            code = next((w['text'] for w in words if w['top'] < 90 and ROUNDCODE.match(w['text'])), None)
            head_w = next((w for w in words if w['text'].endswith('リザルト') and w['top'] < 400), None)  # 題名が長い年は見出しが下がる
            heading = None
            if head_w:
                heading = ' '.join(w['text'] for w in sorted(words, key=lambda w: w['x0'])
                                   if abs(w['top'] - head_w['top']) < 2)
            pace = page_pace(words)
            marks = sorted([w for w in words if re.fullmatch(r'【(男子|女子)】', w['text'])], key=lambda w: w['top'])
            heads = sorted([w for w in words if w['text'] == '順位'], key=lambda w: w['top'])
            codex = sorted(w['top'] for w in words if w['text'] == 'CODEX')
            info = dict(page=pno, tables=0, stray=[])
            for k, hdr in enumerate(heads):
                nxt = heads[k + 1]['top'] if k + 1 < len(heads) else page.height
                end = min([t for t in codex if t > hdr['top']] + [nxt, page.height])
                mark = [m for m in marks if m['top'] < hdr['top']]
                if mark:
                    gender = mark[-1]['text'].strip('【】')
                else:
                    gender = next((g for g in GENDER_CODE if heading and heading.startswith(g)), None)
                if gender is None and heading and heading.startswith('男女'):
                    # 男女合同ページ: 左上の記号（F-w/m = 女子→男子）の順に表が並ぶ
                    order = mixed_order(code)
                    gender = order[k] if k < len(order) else None
                # 見出しの語が文字間隔で割れて印字される年がある（2017 は 'Time' が 'Tim' 'e'）。
                # 隣り合う断片をつないでから列を決める。
                hrow = merge_fragments(sorted([w for w in words if abs(w['top'] - hdr['top']) < 2], key=lambda w: w['x0']))
                nturn, columns = column_map(hrow, [w for w in words if hdr['top'] - 10 <= w['top'] < hdr['top'] - 1])
                bibx = next(w['x0'] for w in hrow if w['text'] == 'BIB')
                lines = _lines([w for w in words if hdr['top'] + 5 < w['top'] < end - 1
                                and not re.fullmatch(r'【(男子|女子)】', w['text'])])

                def bib_word(line):
                    # BIB は 3 桁まで。2 行目の 7 桁 FIS 番号が BIB 列の近くに印字される様式（ばんけい 2026）を選手行と取り違えない
                    return next((w for w in line['words'] if bibx - 6 <= w['x0'] <= bibx + 16 and w['text'].isdigit() and len(w['text']) <= 3), None)

                anchors = [i for i, l in enumerate(lines) if bib_word(l)]
                if anchors and anchors[0]:
                    info['stray'].append(anchors[0])
                blocks = []
                for j, i in enumerate(anchors):
                    blk = lines[i:(anchors[j + 1] if j + 1 < len(anchors) else len(lines))]
                    blocks.append(dict(lines=blk, tokens=[w['text'] for l in blk for w in l['words']],
                                       bib=bib_word(blk[0])['text'], page=pno))
                tables.append(dict(page=pno, code=code, heading=heading, gender=gender,
                                   pace=pace.get(gender, pace.get(None)), nturn=nturn, columns=columns, blocks=blocks))
                info['tables'] += 1
            pages.append(info)
    return year, pages, tables


def group_rounds(tables):
    """連続する表で (見出し, ラウンド記号, 性別) が同じものを1ラウンドにまとめ、
    性別ごとに人数の多い順＝早いラウンドとしてラウンド記号を決める（印字の見出し語は使わない）。"""
    rounds = []
    for t in tables:
        if t['heading'] is None and t['gender'] is None and rounds:
            # 見出し語の無いページ（表が次ページに続くとき）: 直前の表の続きとみなす
            t['heading'], t['code'], t['gender'] = rounds[-1]['heading'], rounds[-1]['printed_code'], rounds[-1]['gender']
            t['pace'] = t['pace'] if t['pace'] is not None else rounds[-1]['tables'][0]['pace']
        key = (t['heading'], t['code'], t['gender'])
        if rounds and rounds[-1]['key'] == key:
            rounds[-1]['tables'].append(t)
        else:
            rounds.append(dict(key=key, heading=t['heading'], printed_code=t['code'], gender=t['gender'], tables=[t]))
    for r in rounds:
        r['blocks'] = [b for t in r['tables'] for b in t['blocks']]
        r['nturn'] = r['tables'][0]['nturn']
        r['columns'] = r['tables'][0]['columns']
        r['pace'] = r['tables'][0]['pace']
        r['pages'] = sorted({t['page'] for t in r['tables']})
    problems = []
    # 同じ表が2回印字されている PDF（結合ミス）: 性別・記号・BIB の集合が同じラウンドは後の方を捨てる
    seen = set()
    kept = []
    for r in rounds:
        sig = (r['gender'], r['printed_code'], tuple(sorted(b['bib'] for b in r['blocks'])))
        if sig in seen:
            continue
        seen.add(sig)
        kept.append(r)
    rounds[:] = kept
    by_gender = collections.defaultdict(list)
    for r in rounds:
        by_gender[r['gender']].append(r)
    most = max((len(rs) for rs in by_gender.values()), default=0)
    full = CODES_BY_ROUNDS.get(most)
    for g, rs in by_gender.items():
        counts = [len(r['blocks']) for r in rs]
        printed = [(r['printed_code'] or '').split('-')[0] for r in rs]
        if g in GENDER_CODE and all(printed) and len(set(printed)) == len(printed) and set(printed) <= {'Q', 'F', 'SF'}:
            # 左上の印字記号（Q-m / F-m / SF-m）が揃っていればそれを使う（人数が同じラウンドがある大会でも決まる）
            for r, c in zip(rs, printed):
                r['code'] = c
            continue
        if full is None or g not in GENDER_CODE or len(set(counts)) != len(counts):
            problems.append('性別 %r のラウンドに記号を決められない（人数 %s）' % (g, counts))
            for r in rs:
                r['code'] = None
            continue
        # ラウンド数が少ない性別は前のラウンドが PDF にない。最終ラウンド側に揃える（2019 女子: SF/F）。
        for r, c in zip(sorted(rs, key=lambda r: -len(r['blocks'])), full[len(full) - len(rs):]):
            r['code'] = c
    return rounds, problems


# ---------------------------------------------------------------- xlsx 側

def read_xlsx(path):
    wb = openpyxl.load_workbook(path)
    sheets, meta = {}, []
    for ws in wb.worksheets:
        hdr = [c.value for c in ws[1]]
        rows = [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
        if ws.title == '_meta':
            meta = rows
        elif ws.title not in NON_ROUND_SHEETS:
            sheets[ws.title] = dict(title=ws.title, hdr=hdr, rows=rows)
    return sheets, meta, [ws.title for ws in wb.worksheets]


def xlsx_numbers(r, nturn):
    keys = (['Rank', 'Bib', 'SAJ No', 'FIS No']
            + ['J%d Base' % i for i in range(1, nturn + 1)] + ['J%d Ded' % i for i in range(1, nturn + 1)]
            + ['Turns Total', 'Jump1', 'DD1', 'Air1 Ja', 'Air1 Jb', 'Jump2', 'DD2', 'Air2 Ja', 'Air2 Jb',
               'Air Total', 'Time', 'Time Points', 'Score', 'Tie'])
    c = collections.Counter()
    for k in keys:
        v = r.get(k)
        if v is None or v == '' or (isinstance(v, str) and not NUM.match(v)):
            continue
        c[D(str(v)).normalize()] += 1
    return c


def printed_numbers(tokens):
    return collections.Counter(D(t).normalize() for t in tokens if NUM.match(t))


# ---------------------------------------------------------------- 各層

def check_completeness(pdf_path, year, pages, rounds, problems, sheets, meta, E, W, N):
    with pdfplumber.open(pdf_path) as pdf:
        npages = len(pdf.pages)
    N['完全性'] += 1
    if len(pages) != npages:
        E['完全性'].append('処理ページ数 %d != PDF %d' % (len(pages), npages))
    for p in pages:
        N['完全性'] += 1
        if not p['tables']:
            E['完全性'].append('p%d: 表が見つからない' % p['page'])
        if p['stray']:
            E['完全性'].append('p%d: 最初の選手より上に表の行 %s（前ページから割れた選手？）' % (p['page'], p['stray']))
    for pr in problems:
        E['完全性'].append(pr)

    expected = {}
    for r in rounds:
        if r['code'] is None or r['gender'] not in GENDER_CODE:
            continue
        r['sheet'] = '%s_%s-%s' % (year, r['code'], GENDER_CODE[r['gender']])
        expected[r['sheet']] = r
        pc = (r['printed_code'] or '').split('-')[0]
        if pc != r['code']:
            W.append('%s: 人数と順序で決めた記号 %s と印字の記号 %s が違う' % (r['sheet'], r['code'], r['printed_code']))
    N['完全性'] += 1
    if set(expected) != set(sheets):
        E['完全性'].append('シート名 PDF から決めたもの %s != xlsx %s' % (sorted(expected), sorted(sheets)))

    for name, r in expected.items():
        sh = sheets.get(name)
        if sh is None:
            continue
        N['完全性'] += 1
        if len(r['blocks']) != len(sh['rows']):
            E['完全性'].append('%s: 行数 xlsx %d != PDF 選手数 %d' % (name, len(sh['rows']), len(r['blocks'])))
        for b in r['blocks']:
            N['完全性'] += 1
            if len(b['lines']) != 2:
                E['完全性'].append('%s BIB %s: 選手ブロックが %d 行（1選手=2行の崩れ）' % (name, b['bib'], len(b['lines'])))
        ok = [x for x in sh['rows'] if x['Status'] == 'OK']
        for i, x in enumerate(ok):
            N['完全性'] += 1
            tied_with_prev = i > 0 and x['Rank'] == ok[i - 1]['Rank'] and x['Score'] == ok[i - 1]['Score']
            if x['Rank'] != i + 1 and not tied_with_prev:
                E['完全性'].append('%s: 順位の連続性 %d 番目の Rank=%s' % (name, i + 1, x['Rank']))
        N['完全性'] += 1
        if any(x['Status'] == 'OK' for x in sh['rows'][len(ok):]):
            E['完全性'].append('%s: 完走者と DNF 等の並びが崩れている' % name)

    # 後のラウンドの選手が前のラウンドにいる（性別ごと）
    for g in GENDER_CODE:
        rs = sorted([r for r in expected.values() if r['gender'] == g], key=lambda r: ROUND_ORDER[r['code']])
        for a, b in zip(rs, rs[1:]):
            if a['sheet'] not in sheets or b['sheet'] not in sheets:
                continue
            prev = {x['FIS No'] for x in sheets[a['sheet']]['rows']}
            for x in sheets[b['sheet']]['rows']:
                N['完全性'] += 1
                if x['FIS No'] not in prev:
                    E['完全性'].append('%s の %s が %s にいない' % (b['sheet'], x['Name'], a['sheet']))

    # _meta の対応表
    by_sheet = {m.get('シート名'): m for m in meta}
    N['完全性'] += 1
    if set(by_sheet) != set(expected):
        E['完全性'].append('_meta のシート名 %s != %s' % (sorted(k for k in by_sheet if k), sorted(expected)))
    for name, r in expected.items():
        m = by_sheet.get(name)
        if not m:
            continue
        N['完全性'] += 1
        want = {'年': int(year), 'ラウンド記号': r['code'], '人数': len(r['blocks']), '印字の見出し語': r['heading'],
                '性別': r['gender'], '印字のラウンド記号': r['printed_code']}
        for k, v in want.items():
            if m.get(k) != v:
                E['完全性'].append('_meta %s %s: %r != PDF %r' % (name, k, m.get(k), v))
        if not m.get('出典URL'):
            E['完全性'].append('_meta %s: 出典URL が空' % name)
    return expected


def check_positions(block, r, columns, tag, E, N):
    """ブロックの各語を最も近いヘッダ列に割り当て、xlsx の同じ列の値と1つずつ比べる。"""
    if len(block['lines']) != 2 or not columns:
        return                                  # 行構造の崩れは完全性の層で止める
    left = columns[0][0] - 15
    printed = {}
    for li, line in enumerate(block['lines']):
        for w in line['words']:
            if w['x0'] < left or w['text'] in STATUS:
                continue
            x, name1, name2 = min(columns, key=lambda c: abs(c[0] - w['x0']))
            key = name1 if li == 0 else name2
            if key is None or key in printed:
                E['印字照合'].append('%s: 列の割り当てができない語 %r（x0=%.0f）' % (tag, w['text'], w['x0']))
                continue
            printed[key] = w['text']
    for key in sorted({c[1] for c in columns} | {c[2] for c in columns if c[2]}):
        N['印字照合'] += 1
        p, x = printed.get(key), r.get(key)
        if p is None and x in (None, ''):
            continue
        if r['Status'] != 'OK' and key == 'Time Points' and p is not None and D(p) == 0 and x is None:
            continue                            # DNF の 0.00 は多重集合の照合で警告に出している
        if p is None or x is None:
            same = False
        elif NUM.match(p) and (isinstance(x, (int, float)) or NUM.match(str(x))):
            same = D(str(x)) == D(p)
        else:
            same = str(x) == p
        if not same:
            E['印字照合'].append('%s: 列 %s 印字 %r != xlsx %r' % (tag, key, p, x))


def check_tokens(expected, sheets, E, W, N):
    """選手ブロックの印字数値の多重集合と xlsx の行の数値を両方向で比べ、列の位置でも1つずつ比べる。"""
    for name, rnd in expected.items():
        sh = sheets.get(name)
        if sh is None:
            continue
        for b, r in zip(rnd['blocks'], sh['rows']):
            N['印字照合'] += 1
            tag = '%s BIB %s' % (name, b['bib'])
            if str(r['Bib']) != b['bib']:
                E['印字照合'].append('%s: xlsx の行の BIB が %s（行の取り違え）' % (tag, r['Bib']))
                continue
            want, got = xlsx_numbers(r, rnd['nturn']), printed_numbers(b['tokens'])
            only_x, only_p = want - got, got - want
            if r['Status'] != 'OK' and not only_x and only_p == collections.Counter({D('0'): 1}):
                W.append('%s %s（%s）: 印字の得点欄 0.00 は xlsx では空欄' % (tag, r['Name'], r['Status']))
            elif only_x or only_p:
                E['印字照合'].append('%s %s: xlsxにだけある数値 %s / 印字にだけある数値 %s'
                                  % (tag, r['Name'], {str(k): v for k, v in only_x.items()},
                                     {str(k): v for k, v in only_p.items()}))
            check_positions(b, r, rnd['columns'], tag, E, N)
            joined = ''.join(b['tokens'])
            for k in ('Name', 'SAJ No', 'FIS No', 'Jump1', 'Jump2', '所属', 'Club', 'Status'):
                v = r.get(k)
                if v in (None, '') or (k == 'Status' and v == 'OK'):
                    continue
                N['印字照合'] += 1
                if str(v).replace(' ', '') not in joined:
                    E['印字照合'].append('%s: %s=%r が印字にない' % (tag, k, v))


def turns_total(base, ded, rule):
    """ターン点。trim=True はベース点・減点それぞれ最高と最低を除いた和の差。下限は floor。"""
    if rule['trim']:
        b = sum(base) - max(base) - min(base)
        d = sum(ded) - max(ded) - min(ded)
    else:
        b, d = sum(base), sum(ded)
    return max(b - d, D(rule['floor']))


def recompute_row(r, nturn, pace, rules):
    """その年の規則で Decimal 再計算。返り値: 不一致だけの [(項目, 再計算, xlsx の値)]。"""
    base = [dec(r['J%d Base' % i]) for i in range(1, nturn + 1)]
    ded = [dec(r['J%d Ded' % i]) for i in range(1, nturn + 1)]
    if None in base or None in ded:
        return [('ジャッジ点', '空欄あり', None)]
    turns = turns_total(base, ded, rules['turns'][str(nturn)])
    cap = D(rules['air']['cap'])
    air = D(0)
    for k in ('1', '2'):
        ddv, ja, jb = dec(r['DD' + k]), dec(r['Air%s Ja' % k]), dec(r['Air%s Jb' % k])
        air += (min(cap, trunc2(ja * ddv)) + min(cap, trunc2(jb * ddv))) / 2
    air = trunc2(air)
    t = rules['time']
    tp = min(max(trunc2(D(t['a']) - D(t['b']) * dec(r['Time']) / pace), D(0)), D(t['max']))
    score = dec(r['Turns Total']) + dec(r['Air Total']) + dec(r['Time Points'])
    return [(name, calc, got) for name, calc, got in (
        ('Turns Total', turns, r['Turns Total']), ('Air Total', air, r['Air Total']),
        ('Time Points', tp, r['Time Points']), ('Score', score, r['Score'])) if calc != dec(got)]


def check_recompute(expected, sheets, rules, E, W, N):
    for name, rnd in expected.items():
        sh = sheets.get(name)
        if sh is None:
            continue
        paces = {t['pace'] for t in rnd['tables']}
        override = rules.get('pace_by_sheet', {}).get(name)
        if override:
            # 印字のペースタイムが無い・性別が違って見える表は、規則ファイルに根拠つきで書いた値を使う
            # （全日本2024 SF: 印字の検定結果は 規則_2024.md）。印字と違えば必ず警告に出す。
            if paces != {D(override['pace'])}:
                W.append('%s: ペースタイムは規則ファイルの %s を使用（印字から読んだ値 %s）。根拠: %s'
                         % (name, override['pace'], sorted(str(p) for p in paces), override['basis']))
            rnd['pace'] = D(override['pace'])
        elif len(paces) != 1 or None in paces:
            E['再計算'].append('%s: ペースタイムが読めない、またはページで違う %s' % (name, paces))
            continue
        if str(rnd['nturn']) not in rules['turns']:
            E['再計算'].append('%s: ターン %d 人の規則が規則ファイルにない' % (name, rnd['nturn']))
            continue
        for r in sh['rows']:
            if r['Status'] != 'OK':
                continue
            N['再計算'] += 1
            tag = '%s %s' % (name, r['Name'])
            try:
                for item, calc, got in recompute_row(r, rnd['nturn'], rnd['pace'], rules):
                    exc = find_exception(rules, name, r['Bib'], item, got, calc)
                    if exc:
                        exc['_used'] = True
                        W.append('%s: %s 印字 %s（再計算 %s）— 規則ファイルの例外（印字を正とする）: %s'
                                 % (tag, item, got, calc, exc['basis']))
                    else:
                        E['再計算'].append('%s: %s 再計算 %s != %s' % (tag, item, calc, got))
            except NotANumber as e:
                E['再計算'].append('%s: 数値でないセル %s' % (tag, e))
    # 例外が使われなかった＝xlsx か規則ファイルのどちらかが変わった。黙って残さない。
    for exc in rules.get('recompute_exceptions', []):
        if not exc.pop('_used', False):
            E['再計算'].append('規則ファイルの例外が一致する行に当たらない: %s BIB %s %s' % (exc['sheet'], exc['bib'], exc['item']))


def find_exception(rules, sheet, bib, item, printed, calc):
    """再計算と印字の食い違いのうち、規則ファイルに根拠つきで登録した1行だけを例外として扱う。
    シート・BIB・項目・印字値・再計算値が全部一致したときだけ。許容誤差は広げない
    （同じ行でも値が変われば計算値が変わるので、例外は効かずエラーになる）。"""
    for exc in rules.get('recompute_exceptions', []):
        if (exc['sheet'], int(exc['bib']), exc['item']) == (sheet, bib, item) \
                and dec(exc['printed']) == dec(printed) and D(exc['calc']) == calc:
            return exc
    return None


def check_ranks(sheets, E, W, N):
    for name, sh in sheets.items():
        ok = [r for r in sh['rows'] if r['Status'] == 'OK']
        try:
            scores = [dec(o['Score']) for o in ok]
        except NotANumber as e:
            E['順位'].append('%s: 数値でない Score %s' % (name, e))
            continue
        for r, sc in zip(ok, scores):
            N['順位'] += 1
            ties = [o['Name'] for o, s2 in zip(ok, scores) if o is not r and s2 == sc]
            calc = 1 + sum(1 for s2 in scores if s2 > sc)
            if ties:
                W.append('%s %s: 同点 %s（同点の決め方は未確認。印字順位 %s）' % (name, r['Name'], ties, r['Rank']))
            elif calc != r['Rank']:
                E['順位'].append('%s %s: 再構成 %d != 印字 %s' % (name, r['Name'], calc, r['Rank']))


def check_cells(expected, sheets, E, N):
    """各ジャッジのベース点・減点・エア点が1セル1値の数値か。列の数が印字の人数と合うか。ラウンド記号列。"""
    for name, rnd in expected.items():
        sh = sheets.get(name)
        if sh is None:
            continue
        n = rnd['nturn']
        cols = (['J%d Base' % i for i in range(1, n + 1)] + ['J%d Ded' % i for i in range(1, n + 1)]
                + ['Air1 Ja', 'Air1 Jb', 'Air2 Ja', 'Air2 Jb', 'DD1', 'DD2'])
        N['セル'] += 1
        missing = [c for c in cols + ['ラウンド記号'] if c not in sh['hdr']]
        extra = [h for h in sh['hdr'] if re.fullmatch(r'J\d+ (Base|Ded)', str(h)) and h not in cols]
        if missing or extra:
            E['セル'].append('%s: 列が印字の人数(%d)と合わない 不足%s 余分%s' % (name, n, missing, extra))
            continue
        for r in sh['rows']:
            N['セル'] += 1
            if r['ラウンド記号'] != rnd['code']:
                E['セル'].append('%s %s: ラウンド記号 %r != %s' % (name, r['Name'], r['ラウンド記号'], rnd['code']))
            if r['Status'] != 'OK':
                continue
            for c in cols:
                N['セル'] += 1
                if not isinstance(r[c], (int, float)) or isinstance(r[c], bool):
                    E['セル'].append('%s %s: %s が数値1つでない (%r)' % (name, r['Name'], c, r[c]))


def check_golden(sheets, golden, E, N):
    index = {(name, r['Bib']): r for name, sh in sheets.items() for r in sh['rows']}
    for g in golden['rows']:
        N['ゴールデン'] += 1
        r = index.get((g['sheet'], g['bib']))
        if r is None:
            E['ゴールデン'].append('%s BIB %s: xlsx に行がない（%s）' % (g['sheet'], g['bib'], g['why']))
            continue
        for k, v in g['values'].items():
            x = r.get(k)
            try:
                numeric = isinstance(v, str) and NUM.match(v) and x is not None
                same = dec(x) == dec(v) if numeric else str(x) == str(v)
            except NotANumber:
                same = False
            if not same:
                E['ゴールデン'].append('%s BIB %s %s: 目視 %r != xlsx %r（守る点: %s）' % (g['sheet'], g['bib'], k, v, x, g['why']))


def verify(pdf_path, xlsx_path, rules, golden=None):
    E, W, N = collections.defaultdict(list), [], collections.Counter()
    year, pages, tables = read_pdf(pdf_path)
    rounds, problems = group_rounds(tables)
    sheets, meta, titles = read_xlsx(xlsx_path)

    expected = check_completeness(pdf_path, year, pages, rounds, problems, sheets, meta, E, W, N)
    check_tokens(expected, sheets, E, W, N)
    check_recompute(expected, sheets, rules, E, W, N)
    check_ranks(sheets, E, W, N)
    check_cells(expected, sheets, E, N)
    if golden:
        check_golden(sheets, golden, E, N)

    # 件数ゲート: 検査が0件のまま「OK」になるのを防ぐ
    n_rows = sum(len(sh['rows']) for sh in sheets.values())
    n_ok = sum(1 for sh in sheets.values() for r in sh['rows'] if r['Status'] == 'OK')
    n_printed = sum(len(r['blocks']) for r in rounds)
    for layer, need in (('印字照合', min(n_rows, n_printed)), ('再計算', n_ok), ('順位', n_ok)):
        if need == 0 or N[layer] < need:
            E['件数ゲート'].append('%s の検査件数 %d < 必要 %d' % (layer, N[layer], need))
    if not golden:
        E['件数ゲート'].append('ゴールデンが未指定')
    elif not golden['rows'] or N['ゴールデン'] != len(golden['rows']):
        E['件数ゲート'].append('ゴールデン検査件数 %d != %d' % (N['ゴールデン'], len(golden['rows'])))

    summary = dict(year=year, pages=len(pages), athletes_printed=n_printed, rows=n_rows, ok_runs=n_ok, sheets=titles,
                   rounds=[(r.get('sheet'), r['heading'], r['printed_code'], r['gender'], r['pages'], len(r['blocks']),
                            r['nturn'], str(r['pace'])) for r in rounds])
    return dict(errors={k: v for k, v in E.items() if v}, warnings=W, counts=dict(N), summary=summary)


def report(res):
    s = res['summary']
    print('%s年 ページ %d / 印字の選手 %d / xlsx の行 %d / 完走 %d' % (s['year'], s['pages'], s['athletes_printed'], s['rows'], s['ok_runs']))
    print('  シート:', s['sheets'])
    for r in s['rounds']:
        print('  %-10s 見出し[%s] 印字記号[%s] %s ページ%s 選手%d ターン%d人 ペース%s' % tuple(r))
    for l in LAYERS:
        print('  %-6s 検査 %5d 件 / エラー %d 件' % (l, res['counts'].get(l, 0), len(res['errors'].get(l, []))))
        for e in res['errors'].get(l, [])[:15]:
            print('      -', e)
        if len(res['errors'].get(l, [])) > 15:
            print('      ... ほか %d 件' % (len(res['errors'][l]) - 15))
    print('  警告 %d 件' % len(res['warnings']))
    for w in res['warnings']:
        print('      *', w)
    ok = not res['errors']
    print('RESULT:', 'OK' if ok else 'NG')
    return ok


if __name__ == '__main__':
    rules = json.load(open(sys.argv[3], encoding='utf-8'))
    golden = json.load(open(sys.argv[4], encoding='utf-8')) if len(sys.argv) > 4 else None
    sys.exit(0 if report(verify(sys.argv[1], sys.argv[2], rules, golden)) else 1)
