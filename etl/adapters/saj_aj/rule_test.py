# -*- coding: utf-8 -*-
"""採点規則の仮説検定。各年の規則を「印字からの全件再計算」で決めるための道具。
xlsx の値（＝印字と照合済みの値）で、ターン・エア・タイム点の計算式の候補ごとに、合わない走の数を数える。
ペースタイムは PDF から読んだ値に加え、--pace で候補を足すと表ごとに一致数を出す。

使い方: python rule_test.py <pdf> <xlsx> [--pace 26.66,23.3]
"""
import sys, collections
from decimal import Decimal as D, ROUND_DOWN, ROUND_HALF_UP
import openpyxl
from verify_nc import read_pdf, group_rounds, GENDER_CODE, trunc2


def r2(x):
    return x.quantize(D('0.01'), rounding=ROUND_HALF_UP)


def trim(v):
    return sum(v) - max(v) - min(v)


TURNS = {
    3: {
        '合計の差・下限0.3': lambda b, d: max(sum(b) - sum(d), D('0.3')),
        'ジャッジごと max(差,0.1) の和': lambda b, d: sum(max(x - y, D('0.1')) for x, y in zip(b, d)),
        '合計の差・下限なし': lambda b, d: sum(b) - sum(d),
    },
    5: {
        'ベース・減点それぞれ最高最低除外の差・下限0.3': lambda b, d: max(trim(b) - trim(d), D('0.3')),
        '同上・下限なし': lambda b, d: trim(b) - trim(d),
        'ジャッジごとの差から最高最低除外・下限0.3': lambda b, d: max(trim([x - y for x, y in zip(b, d)]), D('0.3')),
        'ジャッジごと max(差,0.1) から最高最低除外': lambda b, d: trim([max(x - y, D('0.1')) for x, y in zip(b, d)]),
        '5人全員の和の差・下限0.3': lambda b, d: max(sum(b) - sum(d), D('0.3')),
    },
}
AIR = {
    'ジャッジごと trunc2(点×DD) の平均→合計→trunc2': lambda js: trunc2(sum((trunc2(a * x) + trunc2(b * x)) / 2 for x, a, b in js)),
    'ジャンプごと平均を trunc2→合計': lambda js: sum(trunc2((trunc2(a * x) + trunc2(b * x)) / 2) for x, a, b in js),
    'trunc2(2人平均×DD)→合計': lambda js: sum(trunc2((a + b) / 2 * x) for x, a, b in js),
    '丸めずに合計→trunc2': lambda js: trunc2(sum((a * x + b * x) / 2 for x, a, b in js)),
    'round2(2人平均×DD)→合計': lambda js: sum(r2((a + b) / 2 * x) for x, a, b in js),
}
TIME = {
    'trunc2(48−32×秒/ペース) を0〜20': lambda t, p: min(max(trunc2(D(48) - D(32) * t / p), D(0)), D(20)),
    'round2(…) を0〜20': lambda t, p: min(max(r2(D(48) - D(32) * t / p), D(0)), D(20)),
}


def main(pdf, xlsx, extra_paces):
    year, pages, tables = read_pdf(pdf)
    rounds, problems = group_rounds(tables)
    wb = openpyxl.load_workbook(xlsx)
    runs = []
    for r in rounds:
        name = '%s_%s-%s' % (year, r['code'], GENDER_CODE.get(r['gender'], '?'))
        ws = wb[name]
        hdr = [c.value for c in ws[1]]
        for v in ws.iter_rows(min_row=2, values_only=True):
            x = dict(zip(hdr, v))
            if x['Status'] == 'OK':
                x['_sheet'], x['_pace'], x['_n'] = name, r['pace'], r['nturn']
                runs.append(x)
    d = lambda v: D(str(v))
    print('%s年 完走 %d 走 / ラウンド %s' % (year, len(runs), [(('%s_%s' % (r['code'], r['gender'])), r['nturn'], str(r['pace'])) for r in rounds]))

    for n in sorted({x['_n'] for x in runs}):
        sub = [x for x in runs if x['_n'] == n]
        print('\n[ターン %d人] %d 走' % (n, len(sub)))
        for label, f in TURNS[n].items():
            bad = [x for x in sub if f([d(x['J%d Base' % i]) for i in range(1, n + 1)],
                                       [d(x['J%d Ded' % i]) for i in range(1, n + 1)]) != d(x['Turns Total'])]
            print('  %-44s 合わない %3d  %s' % (label, len(bad), [(b['_sheet'], b['Name']) for b in bad[:3]]))
        below = [x for x in sub if d(x['Turns Total']) == D('0.3') or
                 (n == 5 and trim([d(x['J%d Base' % i]) for i in range(1, 6)]) - trim([d(x['J%d Ded' % i]) for i in range(1, 6)]) < D('0.3')) or
                 (n == 3 and sum(d(x['J%d Base' % i]) for i in range(1, 4)) - sum(d(x['J%d Ded' % i]) for i in range(1, 4)) < D('0.3'))]
        print('  下限0.3の判別に使える走（除外後の差が0.3未満 or 印字0.3）: %d  %s' % (len(below), [(b['_sheet'], b['Name'], b['Turns Total']) for b in below]))

    print('\n[エア]')
    for label, f in AIR.items():
        bad = [x for x in runs if f([(d(x['DD' + k]), d(x['Air%s Ja' % k]), d(x['Air%s Jb' % k])) for k in '12']) != d(x['Air Total'])]
        print('  %-44s 合わない %3d  %s' % (label, len(bad), [(b['_sheet'], b['Name']) for b in bad[:3]]))

    print('\n[タイム] 表ごと（ペースタイム候補別の一致数）')
    sheets = collections.OrderedDict()
    for x in runs:
        sheets.setdefault(x['_sheet'], []).append(x)
    for name, sub in sheets.items():
        cands = {str(sub[0]['_pace'])} | set(extra_paces)
        cands.discard('None')
        for label, f in TIME.items():
            res = {p: sum(1 for x in sub if f(d(x['Time']), D(p)) == d(x['Time Points'])) for p in sorted(cands)}
            print('  %-10s 印字ペース %-6s 完走%2d  %-26s %s' % (name, sub[0]['_pace'], len(sub), label, res))

    print('\n[スコア＝ターン＋エア＋タイム] 合わない %d' % sum(
        1 for x in runs if d(x['Turns Total']) + d(x['Air Total']) + d(x['Time Points']) != d(x['Score'])))


if __name__ == '__main__':
    extra = []
    if '--pace' in sys.argv:
        extra = sys.argv[sys.argv.index('--pace') + 1].split(',')
    main(sys.argv[1], sys.argv[2], extra)
